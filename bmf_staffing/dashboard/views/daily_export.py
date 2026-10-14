"""Day-level research export (daily staffing, base coverage, shift detail).

Every download is a zip: the de-identified data plus a separate key file
linking each random staff ID to a name. IDs are generated fresh per download
(see ``staffing_tool.daily_export.Pseudonymizer``), so the key in the same zip
is the only way back to names. Send the data on; keep the key internal.
"""

import csv
import io
import secrets
import zipfile
from datetime import date, timedelta
from urllib.parse import urlencode

from django.contrib import messages
from django.http import Http404, HttpResponse
from django.shortcuts import redirect, render
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func
from staffing_tool.daily_export import (
    BY_BASE_HEADER,
    EXPORT_ROLES,
    KEY_HEADER,
    NOTES,
    SHIFT_HEADER,
    DailyExport,
    load_daily_export,
)
from staffing_tool.db import session_scope
from staffing_tool.fiscal_year import fy_end_date, fy_label_year
from staffing_tool.models import WeeklyPersonShift
from staffing_tool.timeutil import utc_now_iso as _utc_now_iso

from .dashboard_filters import parse_date_param, parse_fy_week1_from_request
from .helpers import DB_PATH, _ensure_db

DEFAULT_WEEKS_BACK = 8
PREVIEW_DAYS = 14
ROLE_CHOICES = [
    ("", "RN, Medic, EMT"),
    ("RN", "RN"),
    ("MEDIC", "Medic"),
    ("EMT", "EMT"),
]
KEY_FILENAME = "KEY_staff_ids_KEEP_INTERNAL.csv"


def _params(request) -> dict:
    today = date.today()
    default_start = today - timedelta(days=DEFAULT_WEEKS_BACK * 7 - 1)
    date_start = parse_date_param(request.GET.get("date_start", ""), default_start)
    date_end = parse_date_param(request.GET.get("date_end", ""), today)
    if date_start > date_end:
        date_start, date_end = default_start, today
    role = (request.GET.get("role") or "").strip().upper()
    if role not in EXPORT_ROLES:
        role = ""
    return {
        "date_start": date_start,
        "date_end": date_end,
        "role": role,
        "today": today,
    }


def _load(p: dict) -> DailyExport:
    return load_daily_export(
        DB_PATH,
        p["date_start"],
        p["date_end"],
        salt=secrets.token_bytes(32),
        roles=(p["role"],) if p["role"] else EXPORT_ROLES,
    )


def _qs(p: dict) -> str:
    parts = {
        "date_start": p["date_start"].isoformat(),
        "date_end": p["date_end"].isoformat(),
    }
    if p["role"]:
        parts["role"] = p["role"]
    return urlencode(parts)


def daily_export_page(request):
    _ensure_db()
    if not DB_PATH:
        messages.error(request, "Database is not configured (STAFFING_DB_PATH).")
        return redirect("home")
    p = _params(request)
    fy_start = parse_fy_week1_from_request(request, p["today"])
    with session_scope(DB_PATH) as session:
        db_min, db_max = session.query(
            func.min(WeeklyPersonShift.shift_date),
            func.max(WeeklyPersonShift.shift_date),
        ).one()
    export = _load(p)
    return render(
        request,
        "dashboard/daily_export.html",
        {
            "date_start": p["date_start"].isoformat(),
            "date_end": p["date_end"].isoformat(),
            "role_filter": p["role"],
            "role_choices": ROLE_CHOICES,
            "export": export,
            "days": len(export.daily),
            "shift_count": len(export.shifts),
            "staff_count": len(export.key),
            "preview_rows": export.daily[:PREVIEW_DAYS],
            "preview_days": PREVIEW_DAYS,
            "daily_header": export.daily_header,
            "db_date_min": db_min,
            "db_date_max": db_max,
            "fy_start": fy_start.isoformat(),
            "fy_end": fy_end_date(fy_start).isoformat(),
            "fy_label": fy_label_year(fy_start),
            "download_qs": _qs(p),
            "key_filename": KEY_FILENAME,
            "notes": NOTES,
        },
    )


def _csv_bytes(header: list[str], rows: list[list]) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(header)
    writer.writerows(rows)
    return out.getvalue().encode("utf-8-sig")


def _notes_text(export: DailyExport) -> str:
    lines = [
        f"Range {export.date_start} to {export.date_end}.",
        f"{export.weeks_with_detail} of {export.weeks_in_range} weeks have "
        "per-person detail.",
        f"Generated (UTC) {_utc_now_iso()}.",
        "",
        *NOTES,
    ]
    return "\n".join(lines) + "\n"


def _xlsx_bytes(export: DailyExport) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Daily"
    sheets = [
        (ws, export.daily_header, export.daily),
        (wb.create_sheet("By base"), BY_BASE_HEADER, export.by_base),
        (wb.create_sheet("Shift detail"), SHIFT_HEADER, export.shifts),
    ]
    for sheet, header, rows in sheets:
        sheet.append(header)
        for row in rows:
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        for col in sheet.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            sheet.column_dimensions[col[0].column_letter].width = min(width + 2, 40)
    notes = wb.create_sheet("Notes")
    notes.append(["Note"])
    notes["A1"].font = Font(bold=True)
    for line in _notes_text(export).splitlines():
        notes.append([line])
    notes.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _zip_response(export: DailyExport, files: dict[str, bytes], stem: str):
    files[KEY_FILENAME] = _csv_bytes(KEY_HEADER, export.key)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in files.items():
            zf.writestr(name, data)
    response = HttpResponse(buf.getvalue(), content_type="application/zip")
    response["Content-Disposition"] = (
        f'attachment; filename="{stem}_{export.date_start}_to_{export.date_end}.zip"'
    )
    return response


def _export_or_404(request) -> DailyExport:
    _ensure_db()
    if not DB_PATH:
        raise Http404("Database is not configured (STAFFING_DB_PATH).")
    return _load(_params(request))


def daily_export_xlsx(request):
    export = _export_or_404(request)
    return _zip_response(
        export,
        {"daily_staffing.xlsx": _xlsx_bytes(export)},
        "daily_staffing_xlsx",
    )


def daily_export_csv(request):
    export = _export_or_404(request)
    return _zip_response(
        export,
        {
            "daily.csv": _csv_bytes(export.daily_header, export.daily),
            "by_base.csv": _csv_bytes(BY_BASE_HEADER, export.by_base),
            "shift_detail.csv": _csv_bytes(SHIFT_HEADER, export.shifts),
            "README.txt": _notes_text(export).encode("utf-8"),
        },
        "daily_staffing_csv",
    )
