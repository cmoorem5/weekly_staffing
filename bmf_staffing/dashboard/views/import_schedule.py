"""Schedule import view."""

import io
import os
from datetime import UTC, datetime

from django.contrib import messages
from django.http import Http404, HttpResponse
from django.shortcuts import redirect, render
from django.urls import reverse
from django.utils.html import format_html
from openpyxl import Workbook
from staffing_tool.db import session_scope
from staffing_tool.manager_names import canonical_manager_name
from staffing_tool.schedule_apply import apply_schedule_workbook
from staffing_tool.schedule_import import (
    detect_schedule_week_starts,
    parse_schedule_workbook,
)
from staffing_tool.schedule_types import ShiftRecord
from staffing_tool.unit_mappings import save_unit_mappings

from .helpers import (
    _SCHEDULE_UPLOAD_PREFIX,
    DB_PATH,
    _cleanup_old_schedule_uploads,
    _ensure_db,
    _is_uploaded_schedule_path,
    _last_sunday,
    _manager_last_names_upper_for_parse,
    _schedule_upload_dir,
    _training_codes_upper_for_parse,
    backup_staffing_db_before_write,
)
from .manager_shifts import write_manager_aoc_sheet, write_manager_line_shift_sheet

# Reject oversized uploads early (schedule workbooks are well under this).
_MAX_UPLOAD_BYTES = 25 * 1024 * 1024


def _manager_preview_rows(
    records: list[ShiftRecord], roster_upper: frozenset[str]
) -> tuple[list[dict[str, object]], list[dict[str, object]], list[dict[str, object]]]:
    """Manager line shifts + AOC days for this parsed week, grouped like the
    Manager shifts report so counts match once the import is applied."""
    line_rows: list[dict[str, object]] = []
    aoc_rows: list[dict[str, object]] = []
    line_totals: dict[str, int] = {}
    aoc_totals: dict[str, int] = {}
    for r in records:
        if not r.is_manager_row:
            continue
        raw_name = (r.person_display or "").strip() or "(unknown)"
        canon = canonical_manager_name(raw_name, roster_upper)
        raw_for_display = raw_name if raw_name != canon else ""
        if r.manager_event_type == "aoc":
            aoc_totals[canon] = aoc_totals.get(canon, 0) + 1
            aoc_rows.append(
                {
                    "shift_date": r.date.isoformat(),
                    "person_display": canon,
                    "raw_person_display": raw_for_display,
                    "role": r.role,
                    "raw_value": r.raw_value,
                    "source_tab": r.source_tab,
                    "source_cell": r.source_cell,
                }
            )
            continue
        if not (
            r.filled and r.role in {"RN", "MEDIC", "EMT"} and (r.base or "").strip()
        ):
            continue
        line_totals[canon] = line_totals.get(canon, 0) + 1
        line_rows.append(
            {
                "shift_date": r.date.isoformat(),
                "person_display": canon,
                "raw_person_display": raw_for_display,
                "role": r.role,
                "base_name": r.base,
                "service_type": r.service_type,
                "day_night": r.day_night,
                "unit_code": r.unit_code,
                "overtime": r.overtime,
                "raw_value": r.raw_value,
                "source_tab": r.source_tab,
                "source_cell": r.source_cell,
            }
        )
    line_rows.sort(
        key=lambda row: (row["shift_date"], row["person_display"], row["role"])
    )
    aoc_rows.sort(key=lambda row: (row["shift_date"], row["person_display"]))
    summary_rows = sorted(
        (
            {
                "name": name,
                "count": line_totals.get(name, 0),
                "aoc_count": aoc_totals.get(name, 0),
            }
            for name in set(line_totals) | set(aoc_totals)
        ),
        key=lambda row: (
            -(row["count"] + row["aoc_count"]),
            str(row["name"]).lower(),
        ),
    )
    return line_rows, aoc_rows, summary_rows


def _build_import_preview_context(
    upload_path: str, week_start_hint: str
) -> dict | None:
    """Context for import preview, or None if no week headers were found in the file."""
    detected = detect_schedule_week_starts(upload_path)
    if not detected:
        return None
    ws_show = week_start_hint if week_start_hint in detected else detected[0]
    mgr_names = _manager_last_names_upper_for_parse()
    records, issues, _ = parse_schedule_workbook(
        upload_path,
        week_start=ws_show,
        manager_last_names_upper=mgr_names,
        extra_training_codes=_training_codes_upper_for_parse(),
    )
    filled_count = sum(1 for r in records if r.filled)
    ot_count = sum(1 for r in records if r.overtime and r.filled)
    leave_count = sum(1 for r in records if r.leave_type)
    manager_shift_rows, manager_aoc_rows, manager_summary_rows = _manager_preview_rows(
        records, mgr_names
    )
    manager_line_count = len(manager_shift_rows)
    unknown_units = [i for i in issues if i.issue_type == "unknown_unit"]
    return {
        "week_start": ws_show,
        "detected_weeks": detected,
        "upload_path": upload_path,
        "records_count": len(records),
        "filled_count": filled_count,
        "ot_count": ot_count,
        "leave_count": leave_count,
        "manager_line_count": manager_line_count,
        "manager_shift_rows": manager_shift_rows,
        "manager_aoc_rows": manager_aoc_rows,
        "manager_summary_rows": manager_summary_rows,
        "issues": issues,
        "unknown_units": unknown_units,
    }


def import_schedule_manager_export_xlsx(request):
    """Export this previewed (not-yet-applied) week's manager line shifts + AOC days."""
    _ensure_db()
    upload_path = request.GET.get("upload_path", "")
    week_start = (request.GET.get("week_start") or "").strip()
    if not _is_uploaded_schedule_path(upload_path):
        raise Http404("Uploaded file not found on server; please upload again.")
    detected = detect_schedule_week_starts(upload_path)
    if not detected:
        raise Http404("Could not read week dates from this workbook.")
    ws_show = week_start if week_start in detected else detected[0]
    mgr_names = _manager_last_names_upper_for_parse()
    records, _issues, _ = parse_schedule_workbook(
        upload_path,
        week_start=ws_show,
        manager_last_names_upper=mgr_names,
        extra_training_codes=_training_codes_upper_for_parse(),
    )
    line_rows, aoc_rows, summary_rows = _manager_preview_rows(records, mgr_names)

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.append(["Week start", ws_show])
    ws_summary.append(["Note", "Preview only — not yet applied to the database"])
    ws_summary.append([])
    ws_summary.append(["Manager (last name)", "Line shifts", "AOC days"])
    for row in summary_rows:
        ws_summary.append([row["name"], row["count"], row["aoc_count"]])
    ws_summary.append(
        [
            "Total",
            sum(row["count"] for row in summary_rows),
            sum(row["aoc_count"] for row in summary_rows),
        ]
    )

    ws_detail = wb.create_sheet("Line shifts")
    write_manager_line_shift_sheet(ws_detail, line_rows, include_week_start=False)

    ws_aoc = wb.create_sheet("AOC detail")
    write_manager_aoc_sheet(ws_aoc, aoc_rows, include_week_start=False)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    filename = f"manager_shifts_preview_{ws_show}.xlsx"
    response = HttpResponse(
        out.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def import_schedule(request):
    """Upload a schedule workbook, preview parsed results, and optionally apply to create/update a week."""
    _ensure_db()
    _cleanup_old_schedule_uploads(_schedule_upload_dir())

    if request.method == "POST":
        week_start = (request.POST.get("week_start") or _last_sunday()).strip()

        # Re-parse preview after changing which week is selected (same uploaded file).
        if request.POST.get("action") == "update_preview" and request.POST.get(
            "upload_path"
        ):
            upload_path = request.POST.get("upload_path", "")
            if not _is_uploaded_schedule_path(upload_path):
                messages.error(
                    request,
                    "Uploaded file not found on server; please upload again.",
                )
                return redirect("import_schedule")
            week_start = (request.POST.get("week_start") or "").strip()
            ctx = _build_import_preview_context(upload_path, week_start)
            if ctx is None:
                messages.error(
                    request,
                    "Could not read dates from row 1 (columns C–P) on RN & Medic or EMT.",
                )
                return redirect("import_schedule")
            return render(request, "dashboard/import_schedule_preview.html", ctx)

        # Step 2: Apply existing uploaded file to create/update week.
        if request.POST.get("action") == "apply" and request.POST.get("upload_path"):
            upload_path = request.POST.get("upload_path")
            if not _is_uploaded_schedule_path(upload_path):
                messages.error(
                    request, "Uploaded file not found on server; please upload again."
                )
                return redirect("import_schedule")

            detected_apply = detect_schedule_week_starts(upload_path)
            if detected_apply and week_start not in detected_apply:
                messages.error(
                    request,
                    f"Week {week_start} is not in this file. "
                    f"Detected weeks: {', '.join(detected_apply)}.",
                )
                return redirect("import_schedule")

            # Build unit_overrides from mapping form: raw_X -> map_X
            unit_overrides = {}
            for key, map_to in request.POST.items():
                if key.startswith("map_") and map_to and map_to.strip():
                    raw_key = "raw_" + key[4:]
                    raw_val = request.POST.get(raw_key)
                    if raw_val:
                        unit_overrides[raw_val.strip().upper()] = map_to.strip().upper()

            original_name = (
                request.session.pop("schedule_upload_original_name", "") or ""
            )
            request.session.modified = True
            source_filename = original_name or os.path.basename(upload_path)
            entered_by = (
                request.user.username if request.user.is_authenticated else "import"
            )
            # Snapshot the DB first: applying replaces this week's data, so a
            # mis-parsed import stays recoverable from archive/.
            backup_path = backup_staffing_db_before_write()

            roster_added = 0
            with session_scope(DB_PATH) as session:
                if unit_overrides:
                    save_unit_mappings(
                        session,
                        unit_overrides,
                        source="dashboard",
                    )
                result, err = apply_schedule_workbook(
                    session,
                    week_start=week_start,
                    upload_path=upload_path,
                    source_filename=source_filename,
                    unit_overrides=unit_overrides,
                    manager_last_names_upper=_manager_last_names_upper_for_parse(),
                    entered_by=entered_by,
                )
            if err:
                messages.error(request, err)
                return redirect("import_schedule")
            assert result is not None
            roster_added = result.roster_added

            success_msg = f"Week {week_start} imported from schedule."
            if roster_added:
                noun = "staff member" if roster_added == 1 else "staff members"
                success_msg += f" {roster_added} new {noun} added to roster."
            if backup_path:
                success_msg += f" Safety backup saved to archive/{backup_path.name}."
            success_msg += " Please review and export."
            messages.success(request, success_msg)
            if backup_path is not None:
                messages.info(
                    request,
                    format_html(
                        "Safety backup created before this import. "
                        '<a class="alert-link" href="{}">View backups</a>.',
                        reverse("database_backups"),
                    ),
                )
            if backup_path is None:
                messages.warning(
                    request,
                    "Import completed, but the automatic safety backup did not run. "
                    "Use Admin tools → Backup database before your next import.",
                )
            url = reverse("week_edit", kwargs={"week_start": week_start})
            return redirect(f"{url}?imported=1")

        # Step 1: initial upload -> preview.
        file = request.FILES.get("schedule_file")
        if not file:
            messages.error(request, "Please choose a schedule Excel file to upload.")
            return render(
                request,
                "dashboard/import_schedule.html",
                {
                    "week_start": week_start,
                },
            )

        if not (file.name or "").lower().endswith((".xlsx", ".xlsm")):
            messages.error(
                request,
                "Please upload an Excel schedule (.xlsx or .xlsm).",
            )
            return render(
                request,
                "dashboard/import_schedule.html",
                {"week_start": week_start},
            )
        if file.size and file.size > _MAX_UPLOAD_BYTES:
            limit_mb = _MAX_UPLOAD_BYTES // (1024 * 1024)
            messages.error(
                request,
                f"That file is too large (limit {limit_mb} MB). "
                "Please upload the schedule workbook only.",
            )
            return render(
                request,
                "dashboard/import_schedule.html",
                {"week_start": week_start},
            )

        upload_dir = _schedule_upload_dir()
        os.makedirs(upload_dir, exist_ok=True)
        timestamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
        filename = f"{_SCHEDULE_UPLOAD_PREFIX}{timestamp}.xlsx"
        saved_path = os.path.join(upload_dir, filename)
        with open(saved_path, "wb") as dest:
            for chunk in file.chunks():
                dest.write(chunk)

        request.session["schedule_upload_original_name"] = (file.name or "")[:500]
        request.session.modified = True

        ctx = _build_import_preview_context(saved_path, week_start)
        if ctx is None:
            messages.error(
                request,
                "Could not read dates from row 1 (columns C–P) on RN & Medic or EMT. "
                "Check the workbook layout.",
            )
            return render(
                request,
                "dashboard/import_schedule.html",
                {"week_start": week_start},
            )
        detected = ctx["detected_weeks"]
        if week_start not in detected:
            messages.info(
                request,
                f"This file contains week start(s): {', '.join(detected)}. "
                f"Preview below uses {ctx['week_start']}. Choose another week and "
                "click Update preview if needed, then Apply.",
            )
        return render(request, "dashboard/import_schedule_preview.html", ctx)

    return render(
        request,
        "dashboard/import_schedule.html",
        {
            "week_start": _last_sunday(),
        },
    )
