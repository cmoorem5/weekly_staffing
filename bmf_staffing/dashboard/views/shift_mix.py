"""Per-person shift mix and CBA night / weekend requirement report (RN, Medic)."""

import csv
import io
from datetime import date
from urllib.parse import urlencode

from django.contrib import messages
from django.http import Http404, HttpResponse
from django.shortcuts import redirect, render
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func
from staffing_tool.db import session_scope
from staffing_tool.fiscal_year import fy_end_date, fy_label_year
from staffing_tool.models import WeeklyPersonShift
from staffing_tool.person_names import person_sort_key
from staffing_tool.person_ops import list_staff_roster_persons
from staffing_tool.shift_mix import (
    BLOCK_SHIFT_TARGET,
    BLOCK_WEEKS,
    ELIGIBLE_ROLES,
    WEEKEND_SHIFTS_PER_BLOCK,
    WEEKLY_SHIFT_TARGET,
    ShiftMixReport,
    default_block_anchor,
    load_shift_mix,
    night_requirement_table,
)
from staffing_tool.timeutil import utc_now_iso as _utc_now_iso

from .dashboard_filters import parse_date_param, parse_fy_week1_from_request
from .helpers import DB_PATH, _ensure_db

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ROLE_CHOICES = [("", "RN + Medic"), ("RN", "RN"), ("MEDIC", "Medic")]


def _fmt_pct(value: float | None) -> str:
    return "n/a" if value is None else f"{value:g}%"


def _build_context(request) -> dict[str, object]:
    _ensure_db()
    if not DB_PATH:
        raise Http404("Database is not configured (STAFFING_DB_PATH).")

    today = date.today()
    fy_start = parse_fy_week1_from_request(request, today)
    fy_end = fy_end_date(fy_start)
    date_start = parse_date_param(request.GET.get("date_start", ""), fy_start)
    date_end = parse_date_param(request.GET.get("date_end", ""), today)
    if date_start > date_end:
        date_start, date_end = fy_start, today

    role_filter = (request.GET.get("role") or "").strip().upper()
    if role_filter not in ELIGIBLE_ROLES:
        role_filter = ""

    anchor_param = (request.GET.get("block_start") or "").strip()
    default_anchor = default_block_anchor(date_start)
    anchor = parse_date_param(anchor_param, default_anchor)

    roster = [
        (name, role)
        for name, role in list_staff_roster_persons(DB_PATH, role=role_filter or None)
        if role in ELIGIBLE_ROLES
    ]
    roster.sort(key=lambda pair: (person_sort_key(pair[0]), pair[1]))
    groups: dict[str, list[str]] = {}
    for name, _role in roster:
        groups.setdefault(name[:1].upper() or "#", []).append(name)
    person_option_groups = [
        {"letter": letter, "names": names} for letter, names in sorted(groups.items())
    ]
    names = [n for n, _r in roster]
    selected = (request.GET.get("person") or "").strip()
    if not selected and names:
        selected = names[0]
    role_for_person = next((r for n, r in roster if n == selected), role_filter)

    report: ShiftMixReport | None = None
    if selected:
        report = load_shift_mix(
            DB_PATH,
            selected,
            date_start,
            date_end,
            role=role_for_person or None,
            anchor=anchor,
        )

    with session_scope(DB_PATH) as session:
        db_min, db_max = session.query(
            func.min(WeeklyPersonShift.shift_date),
            func.max(WeeklyPersonShift.shift_date),
        ).one()

    qs = urlencode(
        {
            "person": selected,
            "role": role_filter,
            "date_start": date_start.isoformat(),
            "date_end": date_end.isoformat(),
            "block_start": anchor.isoformat(),
        }
    )
    chart = _chart_payload(report) if report else None
    return {
        "report": report,
        "chart": chart,
        "person_option_groups": person_option_groups,
        "person_options": names,
        "selected_person": selected,
        "role_filter": role_filter,
        "role_choices": ROLE_CHOICES,
        "date_start": date_start.isoformat(),
        "date_end": date_end.isoformat(),
        "block_start": anchor.isoformat(),
        "default_block_start": default_anchor.isoformat(),
        "fy_start": fy_start.isoformat(),
        "fy_end": fy_end.isoformat(),
        "fy_label": fy_label_year(fy_start),
        "db_date_min": db_min,
        "db_date_max": db_max,
        "qs": qs,
        "night_table": night_requirement_table(),
        "weekly_target": WEEKLY_SHIFT_TARGET,
        "block_weeks": BLOCK_WEEKS,
        "block_target": BLOCK_SHIFT_TARGET,
        "weekend_target": WEEKEND_SHIFTS_PER_BLOCK,
    }


def _chart_payload(report: ShiftMixReport) -> dict[str, object]:
    return {
        "dayNight": [report.worked.day, report.worked.night],
        "rwGr": [report.worked.rw, report.worked.gr],
        "bases": [
            {
                "label": f"{b.base} ({b.service_type})" if b.service_type else b.base,
                "day": b.day,
                "night": b.night,
            }
            for b in report.by_base
        ],
        "blocks": [
            {
                "label": f"{b.start:%b %-d}",
                "nights": b.nights,
                "requiredNights": b.required_nights,
                "weekend": b.weekend,
                "complete": b.complete,
            }
            for b in report.blocks
        ],
        "weekendTarget": WEEKEND_SHIFTS_PER_BLOCK,
    }


def shift_mix_report(request):
    try:
        ctx = _build_context(request)
    except Http404 as exc:
        messages.error(request, str(exc))
        return redirect("home")
    return render(request, "dashboard/shift_mix.html", ctx)


def _summary_rows(report: ShiftMixReport) -> list[list[object]]:
    hire = report.hire_date.isoformat() if report.hire_date else "missing"
    return [
        ["Person", report.person],
        ["Role", report.role],
        ["Date of hire", hire],
        ["Date start", report.range_start.isoformat()],
        ["Date end", report.range_end.isoformat()],
        ["Block anchor (Sunday)", report.anchor.isoformat()],
        ["Weeks with import data", report.data_weeks],
        ["Weeks without import data", len(report.missing_weeks)],
        ["Worked shifts (excl. OT)", report.worked.total],
        [f"Expected at {WEEKLY_SHIFT_TARGET}/week", report.expected_shifts],
        ["% of weekly target", _fmt_pct(report.target_pct)],
        ["Day shifts", report.worked.day],
        ["Night shifts", report.worked.night],
        ["Day %", _fmt_pct(report.day_pct)],
        ["Night %", _fmt_pct(report.night_pct)],
        ["RW shifts", report.worked.rw],
        ["GR shifts", report.worked.gr],
        ["RW %", _fmt_pct(report.rw_pct)],
        ["GR %", _fmt_pct(report.gr_pct)],
        ["OT shifts (not counted)", report.ot.total],
        ["OT day / night", f"{report.ot.day} / {report.ot.night}"],
        ["OT RW / GR", f"{report.ot.rw} / {report.ot.gr}"],
        ["Training events (not counted)", report.training],
        ["Leave days (not counted)", report.leave_total],
        *[[f"  {k}", v] for k, v in report.leave_counts.items()],
        ["Scored 6-week blocks", len(report.scored_blocks)],
        ["Blocks met", report.blocks_met],
        [
            "Nights worked / required",
            (f"{report.nights_worked_total} / {report.nights_required_total}"),
        ],
        ["Nights %", _fmt_pct(report.nights_pct)],
        [
            "Weekend worked / required",
            (f"{report.weekend_worked_total} / {report.weekend_required_total}"),
        ],
        ["Weekend %", _fmt_pct(report.weekend_pct)],
    ]


_BLOCK_HEADER = [
    "Block start",
    "Block end",
    "Years of service",
    "Worked",
    "Nights",
    "Nights required",
    "Nights +/-",
    "Weekend shifts",
    "Weekend required",
    "Weekend +/-",
    "OT (not counted)",
    "Status",
]


def _block_rows(report: ShiftMixReport) -> list[list[object]]:
    rows = []
    for b in report.blocks:
        rows.append(
            [
                b.start.isoformat(),
                b.end.isoformat(),
                "" if b.years is None else round(b.years, 1),
                b.worked,
                b.nights,
                "" if b.required_nights is None else b.required_nights,
                "" if b.nights_delta is None else b.nights_delta,
                b.weekend,
                WEEKEND_SHIFTS_PER_BLOCK,
                b.weekend_delta,
                b.ot,
                b.status,
            ]
        )
    return rows


_BASE_HEADER = ["Base", "RW/GR", "Day", "Night", "Total"]


def _base_rows(report: ShiftMixReport) -> list[list[object]]:
    return [[b.base, b.service_type, b.day, b.night, b.total] for b in report.by_base]


def _filename(report: ShiftMixReport, ext: str) -> str:
    slug = "".join(c for c in report.person if c.isalnum() or c in "-_ ").strip()
    slug = slug.replace(" ", "_") or "person"
    return (
        f"shift_mix_{slug}_{report.range_start.isoformat()}"
        f"_to_{report.range_end.isoformat()}.{ext}"
    )


def shift_mix_export_csv(request):
    ctx = _build_context(request)
    report = ctx.get("report")
    if not isinstance(report, ShiftMixReport):
        return HttpResponse("No person selected.", status=400)
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["Generated (UTC)", _utc_now_iso()])
    w.writerow(["Report", "Shift mix and CBA requirements"])
    w.writerow([])
    w.writerows(_summary_rows(report))
    w.writerow([])
    w.writerow(["6-week blocks"])
    w.writerow(_BLOCK_HEADER)
    w.writerows(_block_rows(report))
    w.writerow([])
    w.writerow(["By base (worked shifts)"])
    w.writerow(_BASE_HEADER)
    w.writerows(_base_rows(report))
    response = HttpResponse(
        out.getvalue().encode("utf-8-sig"), content_type="text/csv; charset=utf-8"
    )
    response["Content-Disposition"] = (
        f'attachment; filename="{_filename(report, "csv")}"'
    )
    return response


def shift_mix_export_xlsx(request):
    ctx = _build_context(request)
    report = ctx.get("report")
    if not isinstance(report, ShiftMixReport):
        return HttpResponse("No person selected.", status=400)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Measure", "Value"])
    for row in _summary_rows(report):
        ws.append(row)
    ws2 = wb.create_sheet("6-week blocks")
    ws2.append(_BLOCK_HEADER)
    for row in _block_rows(report):
        ws2.append(row)
    ws3 = wb.create_sheet("By base")
    ws3.append(_BASE_HEADER)
    for row in _base_rows(report):
        ws3.append(row)
    ws4 = wb.create_sheet("Notes")
    ws4.append(["Note"])
    for line in (
        "Worked shifts are line shifts that are not overtime. OT, training and "
        "leave are listed separately and never counted toward requirements.",
        "Weekend shifts are Friday night through Sunday night.",
        "Blocks are fixed 6-week periods from the block anchor. Only blocks with "
        "import data for all six weeks that have ended are scored.",
        "Night requirement uses years of service on the block start date.",
        f"Generated (UTC) {_utc_now_iso()}",
    ):
        ws4.append([line])
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for col in sheet.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            sheet.column_dimensions[col[0].column_letter].width = min(width + 2, 70)
    buf = io.BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = (
        f'attachment; filename="{_filename(report, "xlsx")}"'
    )
    return response
