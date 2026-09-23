"""Weekly-grain staffing dashboard export: base coverage + role fill.

Covers the "week" granularity added to buckets_for_range() (staffing_tool)
and the vehicle/base-coverage + role-fill sections added to the staffing
dashboard's CSV/XLSX exports -- a YTD, one-row-per-week export for research
fellows analyzing staffing (staffing rate, OT, exceptions, RW/GR coverage
by base, and role fill), on top of the existing pay-period/month/quarter
aggregation.
"""

import csv
import importlib
import io
import json
import os
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "bmf_staffing"))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "bmf_staffing.settings")

import django

django.setup()

from django.test import Client, RequestFactory
from django.urls import reverse
from openpyxl import load_workbook
from staffing_tool.db import init_db, session_scope
from staffing_tool.metrics import ROLE_FILL_LABELS
from staffing_tool.models import (
    KpiThreshold,
    WeeklyBaseCoverage,
    WeeklyPersonShift,
    WeeklyStaffing,
)
from staffing_tool.time_buckets import buckets_for_range
from tests._temp_db import TempDbTestCase

from dashboard import context_processors
from dashboard.views import helpers

# The `dashboard.views` package re-exports the `staffing_dashboard` VIEW FUNCTION
# under the same name as this module, shadowing `dashboard.views.staffing_dashboard`
# as an attribute -- fetch the actual module via sys.modules to patch its DB_PATH.
staffing_dashboard_view = importlib.import_module("dashboard.views.staffing_dashboard")

WEEK_1 = "2026-07-05"
WEEK_2 = "2026-07-12"


class WeekGranularityBucketsTests(unittest.TestCase):
    def test_buckets_align_to_sunday_and_clip_to_range(self):
        buckets = buckets_for_range("week", date(2026, 7, 1), date(2026, 7, 18))
        # First bucket is clipped at range_start even though the week itself
        # starts on the prior Sunday (2026-06-28).
        self.assertEqual(buckets[0], (date(2026, 7, 1), date(2026, 7, 4)))
        self.assertEqual(buckets[1], (date(2026, 7, 5), date(2026, 7, 11)))
        self.assertEqual(buckets[2], (date(2026, 7, 12), date(2026, 7, 18)))
        self.assertEqual(len(buckets), 3)


class StaffingDashboardWeeklyExportTests(TempDbTestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db_path = str(Path(self.tmp.name) / "week_export.db")
        init_db(self.db_path)
        # Per-week overrides for the "Additional KPIs" section (day/night split,
        # overnights below coverage, pilot vacancies, per-role OT, unpartnered).
        extra_by_week = {
            WEEK_1: {
                "ot_rn_day": 2,
                "ot_rn_night": 1,
                "ot_medic_day": 1,
                "ot_medic_night": 0,
                "ot_emt_day": 0,
                "ot_emt_night": 0,
                "overnights_below": 1,
                "pilot_vacancies": 2,
                "medic_unpartnered": 1,
                "rn_unpartnered_staff": 0,
            },
            WEEK_2: {
                "ot_rn_day": 1,
                "ot_rn_night": 0,
                "ot_medic_day": 0,
                "ot_medic_night": 2,
                "ot_emt_day": 1,
                "ot_emt_night": 1,
                "overnights_below": 0,
                "pilot_vacancies": 0,
                "medic_unpartnered": 0,
                "rn_unpartnered_staff": 2,
            },
        }
        with session_scope(self.db_path) as session:
            for ws in (WEEK_1, WEEK_2):
                session.add(
                    WeeklyStaffing(
                        week_start=ws,
                        day_target=8,
                        night_min=4,
                        filled_day=52,
                        filled_night=30,
                        entered_by="test",
                        created_at=f"{ws}T00:00:00Z",
                        updated_at=f"{ws}T00:00:00Z",
                        **extra_by_week[ws],
                    )
                )
            # Bedford RW: 7/14 (50%) week 1, 14/14 (100%) week 2. No coverage row
            # for other bases -> they report 0% rather than raising.
            session.add(
                WeeklyBaseCoverage(
                    week_start=WEEK_1,
                    base_name="Bedford",
                    rw_staffed_day=7,
                    rw_staffed_night=0,
                    gr_staffed_day=0,
                    gr_staffed_night=0,
                )
            )
            session.add(
                WeeklyBaseCoverage(
                    week_start=WEEK_2,
                    base_name="Bedford",
                    rw_staffed_day=14,
                    rw_staffed_night=0,
                    gr_staffed_day=0,
                    gr_staffed_night=0,
                )
            )
            session.commit()
            # RN role fill: 3 worked week 1 (of 84 capacity), 5 worked week 2.
            for i in range(3):
                session.add(
                    WeeklyPersonShift(
                        week_start=WEEK_1,
                        person_display=f"RN {i}",
                        shift_date=WEEK_1,
                        role="RN",
                        event_type="staffed",
                        included_in_aggregates=1,
                    )
                )
            for i in range(5):
                session.add(
                    WeeklyPersonShift(
                        week_start=WEEK_2,
                        person_display=f"RN {i}",
                        shift_date=WEEK_2,
                        role="RN",
                        event_type="staffed",
                        included_in_aggregates=1,
                    )
                )
            session.commit()
        self._patchers = [
            patch.object(helpers, "DB_PATH", self.db_path),
            patch.object(staffing_dashboard_view, "DB_PATH", self.db_path),
            patch.object(context_processors, "DB_PATH", self.db_path),
        ]
        for p in self._patchers:
            p.start()
        self.client = Client(HTTP_HOST="localhost")

    def _qs(self):
        return {
            "fy": "2026",
            "granularity": "week",
            "date_start": "2026-07-01",
            "date_end": "2026-07-18",
        }

    def test_csv_export_has_one_row_per_week_with_base_and_role_fill_sections(self):
        resp = self.client.get(reverse("staffing_dashboard_export_csv"), self._qs())
        self.assertEqual(resp.status_code, 200)
        text = resp.content.decode("utf-8-sig")
        reader = list(csv.reader(io.StringIO(text)))

        def _section(after_index: int, title: str | None = None):
            """Data rows (keyed by Period) for the section starting at/after ``after_index``.

            Each section is an optional title row, then a "Period,..." column-header
            row, then data rows until a blank line. Returns (rows_by_period, next_index).
            """
            start = after_index
            if title is not None:
                start = next(
                    i
                    for i, r in enumerate(reader)
                    if i >= after_index and r and r[0] == title
                )
            header_idx = next(
                i for i, r in enumerate(reader) if i >= start and r and r[0] == "Period"
            )
            rows: dict[str, list[str]] = {}
            i = header_idx + 1
            while i < len(reader) and reader[i]:
                rows[reader[i][0]] = reader[i]
                i += 1
            return rows, i

        summary_rows, next_idx = _section(0)
        self.assertIn(WEEK_1, summary_rows)
        self.assertIn(WEEK_2, summary_rows)

        base_rows, next_idx = _section(
            next_idx, "Vehicle / base coverage (RW % / GR %, avg across weeks)"
        )
        # Columns after Period/start/end: <Base> RW (%), <Base> GR (%) pairs in
        # BASE_DISPLAY_ORDER (Bedford, Lawrence, Manchester, Mansfield, Plymouth).
        self.assertEqual(base_rows[WEEK_1][3], "50.0")  # Bedford RW %
        self.assertEqual(base_rows[WEEK_1][4], "0.0")  # Bedford GR %
        self.assertEqual(base_rows[WEEK_1][5], "0.0")  # Lawrence RW % (no coverage row)
        self.assertEqual(base_rows[WEEK_2][3], "100.0")  # Bedford RW %

        role_rows, next_idx = _section(
            next_idx, "Role fill — worked vs. seat capacity (pooled)"
        )
        # Columns after Period/start/end: RN worked, RN capacity, RN fill (%), then Medic, EMT.
        self.assertEqual(role_rows[WEEK_1][3], "3")  # RN worked
        self.assertEqual(role_rows[WEEK_1][4], "84")  # RN capacity
        self.assertEqual(role_rows[WEEK_2][3], "5")  # RN worked

    def test_csv_export_has_additional_kpi_section(self):
        resp = self.client.get(reverse("staffing_dashboard_export_csv"), self._qs())
        self.assertEqual(resp.status_code, 200)
        text = resp.content.decode("utf-8-sig")
        reader = list(csv.reader(io.StringIO(text)))
        start = next(i for i, r in enumerate(reader) if r and r[0] == "Additional KPIs")
        header = reader[start + 1]
        rows = {}
        i = start + 2
        while i < len(reader) and reader[i]:
            rows[reader[i][0]] = reader[i]
            i += 1

        def col(name):
            return header.index(name)

        # Day/night staffing rate: filled_day=52/required_day=56, filled_night=30/required_night=28.
        self.assertEqual(rows[WEEK_1][col("Day staffing rate (%, avg)")], "92.86")
        self.assertEqual(rows[WEEK_1][col("Night staffing rate (%, avg)")], "107.14")
        self.assertEqual(rows[WEEK_1][col("Overnights below coverage (total)")], "1")
        self.assertEqual(rows[WEEK_2][col("Overnights below coverage (total)")], "0")
        self.assertEqual(rows[WEEK_1][col("Pilot vacancies (avg)")], "2.0")
        self.assertEqual(rows[WEEK_2][col("Pilot vacancies (avg)")], "0.0")
        self.assertEqual(rows[WEEK_1][col("RN OT (shifts, total)")], "3")  # 2 + 1
        self.assertEqual(rows[WEEK_2][col("Medic OT (shifts, total)")], "2")  # 0 + 2
        self.assertEqual(rows[WEEK_2][col("EMT OT (shifts, total)")], "2")  # 1 + 1
        self.assertEqual(rows[WEEK_1][col("Medic unpartnered (total)")], "1")
        self.assertEqual(rows[WEEK_2][col("RN unpartnered (total)")], "2")

    def test_xlsx_export_has_base_coverage_and_role_fill_sheets(self):
        resp = self.client.get(reverse("staffing_dashboard_export_xlsx"), self._qs())
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(io.BytesIO(resp.content))
        self.assertIn("Vehicle-Base coverage", wb.sheetnames)
        self.assertIn("Role fill", wb.sheetnames)
        self.assertIn("Additional KPIs", wb.sheetnames)

        ws_extra = wb["Additional KPIs"]
        extra_header = [c.value for c in ws_extra[1]]
        extra_rows = {row[0].value: row for row in ws_extra.iter_rows(min_row=2)}
        rn_ot_col = extra_header.index("RN OT (shifts, total)")
        self.assertEqual(extra_rows[WEEK_1][rn_ot_col].value, 3)

        ws_base = wb["Vehicle-Base coverage"]
        header = [c.value for c in ws_base[1]]
        self.assertIn("Bedford RW (%)", header)
        rows = {row[0].value: row for row in ws_base.iter_rows(min_row=2)}
        bedford_rw_col = header.index("Bedford RW (%)")
        self.assertAlmostEqual(rows[WEEK_1][bedford_rw_col].value, 50.0)
        self.assertAlmostEqual(rows[WEEK_2][bedford_rw_col].value, 100.0)

        ws_role = wb["Role fill"]
        role_header = [c.value for c in ws_role[1]]
        rn_label = ROLE_FILL_LABELS["RN"]
        self.assertIn(f"{rn_label} worked", role_header)
        role_rows = {row[0].value: row for row in ws_role.iter_rows(min_row=2)}
        rn_worked_col = role_header.index(f"{rn_label} worked")
        rn_capacity_col = role_header.index(f"{rn_label} capacity")
        self.assertEqual(role_rows[WEEK_1][rn_worked_col].value, 3)
        self.assertEqual(role_rows[WEEK_1][rn_capacity_col].value, 84)
        self.assertEqual(role_rows[WEEK_2][rn_worked_col].value, 5)

    def test_legacy_week_per_role_ot_uses_aggregate_columns(self):
        # A week imported before the day/night OT split stores per-role OT only
        # in ot_rn/ot_medic/ot_emt; the Additional KPIs table must still show it.
        with session_scope(self.db_path) as session:
            row = session.query(WeeklyStaffing).filter_by(week_start=WEEK_2).one()
            for f in (
                "ot_rn_day",
                "ot_rn_night",
                "ot_medic_day",
                "ot_medic_night",
                "ot_emt_day",
                "ot_emt_night",
            ):
                setattr(row, f, 0)
            row.ot_rn, row.ot_medic, row.ot_emt = 4, 2, 1
            session.commit()
        request = RequestFactory().get("/", self._qs())
        ctx = staffing_dashboard_view._build_staffing_dashboard_context(request)
        extra = {r["bucket_start"]: r for r in ctx["extra_kpi_table"]}
        self.assertEqual(
            (
                extra[WEEK_2]["ot_rn_total"],
                extra[WEEK_2]["ot_medic_total"],
                extra[WEEK_2]["ot_emt_total"],
            ),
            (4, 2, 1),
        )

    def test_chart_targets_follow_kpi_threshold_direction(self):
        with session_scope(self.db_path) as session:
            session.query(KpiThreshold).delete()
            session.add(
                KpiThreshold(
                    metric_name="Staffing Rate", green_min=0.95, higher_is_better=1
                )
            )
            session.add(
                KpiThreshold(
                    metric_name="OT Dependency", green_max=0.10, higher_is_better=0
                )
            )
            session.commit()
        request = RequestFactory().get("/", self._qs())
        ctx = staffing_dashboard_view._build_staffing_dashboard_context(request)
        targets = json.loads(ctx["chart_targets_json"])
        # green_min for higher-is-better, green_max for lower-is-better.
        self.assertEqual(targets["staffing_rate"], 95.0)
        self.assertEqual(targets["ot_dependency"], 10.0)
        self.assertEqual(json.loads(ctx["weeks_per_bucket_json"]), [1, 1])


if __name__ == "__main__":
    unittest.main()
