"""Per-role OT, shift-exception totals and role fill agree across report paths.

Legacy weeks (imported before the day/night OT split) carry OT only in the
per-role aggregate columns; every report must read them through the same
fallback ``compute_week_metrics`` uses, and every "shift exceptions" total
must include PFML the way the Shift Exception % KPI does.
"""

import os
import unittest

from openpyxl import load_workbook
from staffing_tool.db import session_scope
from staffing_tool.metrics import (
    compute_role_fill,
    compute_role_fill_by_week,
    compute_week_metrics,
    role_ot_totals,
    weekly_leave_total,
)
from staffing_tool.models import WeeklyPersonShift, WeeklyStaffing
from staffing_tool.monthly_html_report import load_monthly_board_data
from staffing_tool.monthly_report import export_monthly_report
from staffing_tool.report_style import _axis_top
from tests._temp_db import TempDbTestCase


class RoleOtTotalsTests(unittest.TestCase):
    def test_uses_day_night_split_when_present(self):
        row = WeeklyStaffing(
            ot_rn_day=2,
            ot_rn_night=1,
            ot_medic_day=0,
            ot_medic_night=4,
            ot_emt_day=1,
            ot_emt_night=0,
            ot_rn=99,
            ot_medic=99,
            ot_emt=99,
        )
        self.assertEqual(role_ot_totals(row), {"RN": 3, "MEDIC": 4, "EMT": 1})

    def test_falls_back_to_legacy_per_role_columns(self):
        row = WeeklyStaffing(ot_rn=5, ot_medic=2, ot_emt=1)
        self.assertEqual(role_ot_totals(row), {"RN": 5, "MEDIC": 2, "EMT": 1})

    def test_per_role_sum_matches_metrics_total_on_legacy_week(self):
        row = WeeklyStaffing(
            week_start="2025-01-05",
            filled_day=50,
            filled_night=25,
            overnights_below=0,
            pilot_vacancies=0,
            ot_rn=5,
            ot_medic=2,
            ot_emt=1,
        )
        m = compute_week_metrics(row, [], [])
        self.assertEqual(sum(role_ot_totals(row).values()), m.ot_shifts)


class WeeklyLeaveTotalTests(unittest.TestCase):
    def test_includes_pfml_jury_and_brev(self):
        row = WeeklyStaffing(
            week_start="2025-01-05",
            filled_day=50,
            filled_night=25,
            overnights_below=0,
            pilot_vacancies=0,
            leave_at=1,
            leave_lt=2,
            leave_sick=3,
            leave_loa=4,
            leave_pfml=5,
            leave_jury=6,
            leave_brev=7,
        )
        self.assertEqual(weekly_leave_total(row), 28)
        self.assertEqual(compute_week_metrics(row, [], []).leave_total, 28)


class AxisTopTests(unittest.TestCase):
    def test_keeps_floor_for_normal_values(self):
        self.assertEqual(_axis_top([4.0, 12.5, 29.9], floor=30), 30)

    def test_expands_past_floor_so_spikes_stay_on_chart(self):
        self.assertEqual(_axis_top([12.0, 34.2], floor=30), 40)
        self.assertEqual(_axis_top([30.0], floor=30), 40)

    def test_empty_series(self):
        self.assertEqual(_axis_top([], floor=110), 110)


class ReportPathConsistencyTests(TempDbTestCase):
    def setUp(self):
        self.make_temp_db()
        with session_scope(self.db_path) as session:
            # Legacy week: OT only in the per-role aggregate columns, plus PFML.
            session.add(
                WeeklyStaffing(
                    week_start="2025-12-07",
                    filled_day=50,
                    filled_night=20,
                    ot_rn=4,
                    ot_medic=3,
                    ot_emt=2,
                    ot_shifts=9,
                    leave_at=1,
                    leave_pfml=2,
                )
            )
            # Current-format week: day/night split populated.
            session.add(
                WeeklyStaffing(
                    week_start="2025-12-14",
                    filled_day=52,
                    filled_night=22,
                    ot_rn_day=1,
                    ot_rn_night=1,
                    ot_medic_day=1,
                    ot_rn=2,
                    ot_medic=1,
                    ot_emt=0,
                    ot_shifts=3,
                    leave_sick=3,
                )
            )
            session.commit()
            for week, role, n in [
                ("2025-12-07", "RN", 5),
                ("2025-12-07", "EMT", 2),
                ("2025-12-14", "RN", 3),
                ("2025-12-14", "MEDIC", 4),
            ]:
                for i in range(n):
                    session.add(
                        WeeklyPersonShift(
                            week_start=week,
                            person_display=f"{role} {i}",
                            shift_date=week,
                            role=role,
                            event_type="staffed",
                            included_in_aggregates=1,
                            source_tab="RN",
                            source_cell=f"C{i + 2}",
                        )
                    )
            session.commit()

    def test_monthly_html_per_role_ot_includes_legacy_week(self):
        data = load_monthly_board_data(self.db_path, "2025-12-01", "2025-12-31")
        self.assertEqual(dict(data.ot_by_role), {"RN": 6, "Paramedic": 4, "EMT": 2})

    def test_monthly_excel_volumes_match_metrics(self):
        path = export_monthly_report(
            self.db_path,
            "2025-12-01",
            "2025-12-31",
            os.path.join(self.tmp.name, "out"),
        )
        wb = load_workbook(path)
        ws = wb[wb.sheetnames[0]]
        values = {
            ws.cell(r, 1).value: ws.cell(r, 2).value
            for r in range(1, ws.max_row + 1)
            if ws.cell(r, 1).value
        }
        # 1 AT + 2 PFML + 3 SICK: PFML counts, as it does in Shift Exception %.
        self.assertEqual(values["Shift exceptions (total)"], 6)
        self.assertEqual(values["RN OT shifts (total)"], 6)
        self.assertEqual(values["Medic OT shifts (total)"], 4)
        self.assertEqual(values["EMT OT shifts (total)"], 2)

    def test_role_fill_by_week_matches_per_week_calls(self):
        weeks = ["2025-12-07", "2025-12-14"]
        with session_scope(self.db_path) as session:
            batched = compute_role_fill_by_week(session, weeks)
            for ws in weeks:
                single = compute_role_fill(session, [ws])
                self.assertEqual(
                    [(r.role, r.worked, r.capacity) for r in batched[ws]],
                    [(r.role, r.worked, r.capacity) for r in single],
                )
            pooled = {r.role: r for r in compute_role_fill(session, weeks)}
        self.assertEqual(pooled["RN"].worked, 8)
        self.assertEqual(pooled["RN"].capacity, 168)
        self.assertEqual(pooled["MEDIC"].worked, 4)
        self.assertEqual(pooled["EMT"].worked, 2)


if __name__ == "__main__":
    unittest.main()
