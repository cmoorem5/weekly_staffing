"""Daily export endpoints: page renders, CSV/XLSX download, bad params fall back."""

import io
import os
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "bmf_staffing"))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "bmf_staffing.settings")

import django

django.setup()

from django.test import Client
from django.urls import reverse
from openpyxl import load_workbook
from staffing_tool.db import session_scope
from staffing_tool.models import WeeklyPersonShift, WeeklyStaffing
from tests._temp_db import TempDbTestCase

from dashboard import context_processors
from dashboard.views import daily_export, helpers

QS = {"date_start": "2026-07-05", "date_end": "2026-07-11"}


class DailyExportViewTests(TempDbTestCase):
    def setUp(self):
        self.make_temp_db()
        with session_scope(self.db_path) as s:
            s.add(WeeklyStaffing(week_start="2026-07-05", filled_day=0, filled_night=0))
            s.flush()
            s.add(
                WeeklyPersonShift(
                    week_start="2026-07-05",
                    person_display="Doe, Jane",
                    shift_date="2026-07-06",
                    role="RN",
                    event_type="staffed",
                    base_name="Bedford",
                    service_type="RW",
                    day_night="D",
                    unit_code="D7B",
                    source_tab="RN",
                    source_cell="C5",
                )
            )
        self._patchers = [
            patch.object(helpers, "DB_PATH", self.db_path),
            patch.object(daily_export, "DB_PATH", self.db_path),
            patch.object(context_processors, "DB_PATH", self.db_path),
        ]
        for p in self._patchers:
            p.start()
        self.client = Client(HTTP_HOST="localhost")

    def test_page_renders(self):
        resp = self.client.get(reverse("daily_export"), QS)
        self.assertEqual(resp.status_code, 200)
        self.assertIn("Daily data export", resp.content.decode())

    def test_csv_datasets(self):
        for dataset, first in (
            ("daily", "date"),
            ("by_base", "date"),
            ("shifts", "staff_id"),
        ):
            resp = self.client.get(
                reverse("daily_export_csv"), {**QS, "dataset": dataset}
            )
            self.assertEqual(resp.status_code, 200)
            self.assertTrue(resp.content.decode("utf-8-sig").startswith(first))
        shifts = self.client.get(
            reverse("daily_export_csv"), {**QS, "dataset": "shifts"}
        )
        self.assertNotIn("Doe", shifts.content.decode("utf-8-sig"))

    def test_xlsx_has_all_sheets(self):
        resp = self.client.get(reverse("daily_export_xlsx"), QS)
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(io.BytesIO(resp.content))
        self.assertEqual(wb.sheetnames, ["Daily", "By base", "Shift detail", "Notes"])
        self.assertEqual(wb["Daily"].max_row, 8)  # header + 7 days

    def test_bad_dates_fall_back(self):
        resp = self.client.get(
            reverse("daily_export"), {"date_start": "nope", "date_end": "2026-01-01"}
        )
        self.assertEqual(resp.status_code, 200)


if __name__ == "__main__":
    unittest.main()
